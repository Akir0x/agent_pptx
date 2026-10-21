import copy, json, re
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn
from pptx.parts.chart import ChartPart
from pptx.parts.embeddedpackage import EmbeddedXlsxPart
R = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"

def rids(el, m):
    for n in el.iter():
        for a, v in list(n.attrib.items()):
            if a.startswith(R) and v in m: n.set(a, m[v])

def clonar_grafico(pkg, cp):
    novo = ChartPart.load(pkg.next_partname("/ppt/charts/chart%d.xml"), cp.content_type, pkg, cp.blob)
    m = {}
    for rid, rel in cp.rels.items():
        if rel.is_external: m[rid] = novo.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.PACKAGE: m[rid] = novo.relate_to(EmbeddedXlsxPart.new(rel.target_part.blob, pkg), RT.PACKAGE)
        else: m[rid] = novo.relate_to(rel.target_part, rel.reltype)
    rids(novo._element, m)
    return novo

def duplicar(prs, orig):
    s = prs.slides.add_slide(orig.slide_layout)
    arv = s.shapes._spTree
    for el in list(arv)[2:]: arv.remove(el)
    bg = orig._element.cSld.find(qn("p:bg"))
    if bg is not None: s._element.cSld.insert(0, copy.deepcopy(bg))
    m = {}
    for rid, rel in orig.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE): continue
        if rel.is_external: m[rid] = s.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.CHART: m[rid] = s.part.relate_to(clonar_grafico(s.part.package, rel.target_part), RT.CHART)
        else: m[rid] = s.part.relate_to(rel.target_part, rel.reltype)
    for el in list(orig.shapes._spTree)[2:]:
        c = copy.deepcopy(el); rids(c, m); arv.append(c)
    return s

def caixas(shapes):
    for sh in shapes:
        yield sh
        if sh.shape_type == 6: yield from caixas(sh.shapes)

def texto(tf, valor):
    itens = [str(v) for v in valor] if isinstance(valor, list) else str(valor).split("\n")
    body = tf._txBody
    ps = body.findall(qn("a:p")); mod0 = ps[0]; mod1 = ps[1] if len(ps) > 1 else ps[0]
    for p in ps: body.remove(p)
    for i, t in enumerate(itens):
        p = copy.deepcopy(mod0 if i == 0 else mod1)
        rs = p.findall(qn("a:r"))
        for x in rs[1:] + p.findall(qn("a:fld")) + p.findall(qn("a:br")): p.remove(x)
        if rs: rs[0].find(qn("a:t")).text = t
        else:
            r = p.makeelement(qn("a:r"), {}); fim = p.find(qn("a:endParaRPr"))
            if fim is not None:
                rp = copy.deepcopy(fim); rp.tag = qn("a:rPr"); r.append(rp); fim.addprevious(r)
            else: p.append(r)
            r.append(r.makeelement(qn("a:t"), {})); r[-1].text = t
        body.append(p)

def montar(biblioteca, pedido, saida):
    prs = Presentation(biblioteca)
    orig = list(prs.slides)
    avisos = []
    for i, item in enumerate(pedido["slides"], 1):
        n = int(item["slide_biblioteca"])
        s = duplicar(prs, orig[n - 1])
        cx = {c.name: c for c in caixas(s.shapes)}
        sugest = []
        for nome, v in item.get("campos", {}).items():
            sh = cx.get(nome)
            if sh is None or nome.startswith("fixo_"): avisos.append(f"slide {i}: caixa {nome} ignorada"); continue
            if getattr(sh, "has_chart", False) and sh.has_chart:
                d = CategoryChartData(); d.categories = v["categorias"]
                for se in v["series"]: d.add_series(se["nome"], se["valores"])
                sh.chart.replace_data(d)
            elif getattr(sh, "has_table", False) and sh.has_table:
                for r, linha in enumerate(v["linhas"][:len(sh.table.rows)]):
                    for c, t in enumerate(linha[:len(sh.table.columns)]): texto(sh.table.cell(r, c).text_frame, t)
            elif nome.startswith("imagem_") or not sh.has_text_frame: sugest.append(f"Trocar {nome}: {v}")
            else: texto(sh.text_frame, v)
        notas = "\n".join([x for x in [item.get("notas", "")] if x] + sugest)
        if notas: s.notes_slide.notes_text_frame.text = notas
    lst = prs.slides._sldIdLst
    for sid in list(lst)[:len(orig)]:
        prs.part.drop_rel(sid.rId); lst.remove(sid)
    prs.save(saida)
    return avisos
