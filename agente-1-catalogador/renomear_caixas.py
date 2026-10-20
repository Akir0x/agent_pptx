"""
Agente 1 - Renomear as caixas automaticamente.

Lê a aba "caixas" do levantamento.xlsx (colunas slide, nome_atual, nome_novo)
e grava os nomes novos dentro do PowerPoint. Assim ninguém precisa renomear
caixa por caixa no Painel de Seleção.

O arquivo original não é alterado: o resultado sai como biblioteca_final.pptx
(ou o nome passado em --saida). É esse arquivo que o agente 2 usa.

Uso:
  python renomear_caixas.py biblioteca.pptx saida/levantamento.xlsx
  python renomear_caixas.py biblioteca.pptx saida/levantamento.xlsx --saida biblioteca_final.pptx
"""

import argparse
import os
import sys

from openpyxl import load_workbook
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn


def listar_caixas(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from listar_caixas(shape.shapes)


def definir_nome(shape, nome):
    """Grava o nome no cNvPr da caixa (funciona para texto, imagem, gráfico e grupo)."""
    cnvpr = shape._element.find(".//" + qn("p:cNvPr"))
    cnvpr.set("name", nome)


def ler_renomeacoes(caminho_xlsx):
    wb = load_workbook(caminho_xlsx, read_only=True, data_only=True)
    if "caixas" not in wb.sheetnames:
        sys.exit('A planilha precisa ter uma aba chamada "caixas".')
    linhas = list(wb["caixas"].iter_rows(values_only=True))
    cabecalho = [str(c).strip().lower() if c else "" for c in linhas[0]]
    for coluna in ("slide", "nome_atual", "nome_novo"):
        if coluna not in cabecalho:
            sys.exit(f'Coluna "{coluna}" não encontrada na aba "caixas".')
    i_slide, i_atual, i_novo = (cabecalho.index(c) for c in ("slide", "nome_atual", "nome_novo"))

    renomeacoes = {}
    for linha in linhas[1:]:
        if not linha or linha[i_slide] in (None, "") or not linha[i_atual] or not linha[i_novo]:
            continue
        slide = int(str(linha[i_slide]).strip())
        renomeacoes.setdefault(slide, {})[str(linha[i_atual]).strip()] = str(linha[i_novo]).strip()
    return renomeacoes


def main():
    parser = argparse.ArgumentParser(description="Renomeia as caixas a partir do levantamento.")
    parser.add_argument("pptx", help="biblioteca.pptx original")
    parser.add_argument("levantamento", help="levantamento.xlsx preenchido")
    parser.add_argument("--saida", default="biblioteca_final.pptx")
    args = parser.parse_args()

    if os.path.abspath(args.pptx) == os.path.abspath(args.saida):
        sys.exit("Use um nome de saída diferente do arquivo original.")

    prs = Presentation(args.pptx)
    renomeacoes = ler_renomeacoes(args.levantamento)
    feitas, nao_achadas, avisos = 0, [], []

    for n, slide in enumerate(prs.slides, start=1):
        pedido = renomeacoes.get(n, {})
        caixas = list(listar_caixas(slide.shapes))

        # Nomes atuais repetidos no mesmo slide deixam a renomeação ambígua
        nomes = [c.name for c in caixas]
        for nome in set(nomes):
            if nomes.count(nome) > 1 and nome in pedido:
                avisos.append(f"slide {n}: há {nomes.count(nome)} caixas chamadas "
                              f'"{nome}"; só a primeira foi renomeada')

        usados = set()
        for caixa in caixas:
            atual = caixa.name
            if atual in pedido and atual not in usados:
                definir_nome(caixa, pedido[atual])
                usados.add(atual)
                feitas += 1
        nao_achadas += [f"slide {n}: {nome}" for nome in pedido if nome not in usados]

        novos = [pedido[a] for a in usados]
        for nome in set(novos):
            if novos.count(nome) > 1:
                avisos.append(f'slide {n}: o nome novo "{nome}" foi usado mais de uma vez')

    prs.save(args.saida)
    print(f"{feitas} caixas renomeadas. Arquivo salvo em {args.saida}")
    for linha in nao_achadas:
        print(f"  Não encontrada: {linha}")
    for linha in avisos:
        print(f"  Atenção: {linha}")


if __name__ == "__main__":
    main()
