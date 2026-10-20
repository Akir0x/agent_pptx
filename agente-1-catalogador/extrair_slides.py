"""
Agente 1 - Etapa automática de extração.

Lê o biblioteca.pptx e gera, na pasta de saída:
  - miniaturas/NN.png ........ imagem de cada slide (para anexar no agente)
  - para_o_agente/NN.txt ..... lista de caixas de cada slide (para colar no agente)
  - caixas_extraidas.xlsx .... todas as caixas com tamanho, fonte e limite estimado
  - levantamento.xlsx ........ planilha vazia com as abas "moldes" e "caixas",
                               onde se colam as tabelas que o agente devolve

Uso:
  python extrair_slides.py biblioteca.pptx
  python extrair_slides.py biblioteca.pptx --saida saida --sem-imagens
"""

import argparse
import os
import platform
import shutil
import subprocess
import sys
import tempfile

from openpyxl import Workbook
from openpyxl.styles import Font
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

EMU_POR_PT = 12700
FONTE_PADRAO = 18  # usada quando a caixa não informa o tamanho da fonte

COLUNAS_MOLDES = ["slide", "molde", "funcao", "layout", "descricao_visual",
                  "quando_usar", "quando_nao_usar", "cores", "densidade_texto",
                  "tem_imagem"]
COLUNAS_CAIXAS = ["slide", "nome_atual", "nome_novo", "o_que_vai",
                  "max_chars", "linhas"]


# --------------------------------------------------------------------------
# Exportar imagens dos slides
# --------------------------------------------------------------------------

def exportar_pngs_powerpoint(arquivo, pasta_img):
    """Windows com PowerPoint instalado: exporta cada slide em 1280x720."""
    import win32com.client  # pywin32

    app = win32com.client.Dispatch("PowerPoint.Application")
    pres = app.Presentations.Open(arquivo, ReadOnly=True, WithWindow=False)
    try:
        for i, slide in enumerate(pres.Slides, start=1):
            slide.Export(os.path.join(pasta_img, f"{i:02d}.png"), "PNG", 1280, 720)
    finally:
        pres.Close()
        app.Quit()


def exportar_pngs_libreoffice(arquivo, pasta_img):
    """Alternativa sem PowerPoint: LibreOffice gera PDF e pdftoppm vira PNG."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    pdftoppm = shutil.which("pdftoppm")
    if not soffice or not pdftoppm:
        raise RuntimeError("LibreOffice (soffice) e pdftoppm não encontrados.")

    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run([soffice, "--headless", "--convert-to", "pdf",
                        "--outdir", tmp, arquivo],
                       check=True, capture_output=True)
        pdf = os.path.join(tmp, os.path.splitext(os.path.basename(arquivo))[0] + ".pdf")
        prefixo = os.path.join(tmp, "s")
        subprocess.run([pdftoppm, "-png", "-scale-to-x", "1280", "-scale-to-y", "720",
                        pdf, prefixo], check=True, capture_output=True)
        geradas = sorted(f for f in os.listdir(tmp) if f.startswith("s") and f.endswith(".png"))
        for i, nome in enumerate(geradas, start=1):
            shutil.move(os.path.join(tmp, nome), os.path.join(pasta_img, f"{i:02d}.png"))


def exportar_pngs(arquivo, pasta_img):
    if platform.system() == "Windows":
        try:
            exportar_pngs_powerpoint(arquivo, pasta_img)
            return
        except Exception as erro:  # sem PowerPoint ou sem pywin32
            print(f"Aviso: não deu para usar o PowerPoint ({erro}). Tentando LibreOffice...")
    exportar_pngs_libreoffice(arquivo, pasta_img)


# --------------------------------------------------------------------------
# Ler as caixas
# --------------------------------------------------------------------------

def listar_caixas(shapes, grupo=""):
    """Percorre as caixas do slide, abrindo grupos por dentro."""
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from listar_caixas(shape.shapes, grupo=shape.name)
        else:
            yield shape, grupo


def tipo_da_caixa(shape):
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return "imagem"
    if getattr(shape, "has_chart", False) and shape.has_chart:
        return "grafico"
    if getattr(shape, "has_table", False) and shape.has_table:
        return "tabela"
    if shape.has_text_frame:
        # formas vazias (círculos, faixas) são decoração; caixas de texto e
        # placeholders vazios ainda podem receber conteúdo
        if (shape.text_frame.text.strip() or shape.is_placeholder
                or shape.shape_type == MSO_SHAPE_TYPE.TEXT_BOX):
            return "texto"
    return "forma"


def tamanho_fonte(shape):
    """Maior tamanho de fonte encontrado na caixa (ou o padrão)."""
    tamanhos = [r.font.size.pt
                for p in shape.text_frame.paragraphs
                for r in p.runs if r.font.size]
    return max(tamanhos) if tamanhos else FONTE_PADRAO


def estimar_limites(shape, fonte):
    """Estimativa: largura média de letra = metade da fonte; 10% de margem."""
    if not shape.width or not shape.height:
        return "", ""
    largura_pt = shape.width / EMU_POR_PT
    altura_pt = shape.height / EMU_POR_PT
    chars_linha = max(1, int(largura_pt / (fonte * 0.5)))
    linhas = max(1, int(altura_pt / (fonte * 1.2)))
    return int(chars_linha * linhas * 0.9), linhas


def pt(valor):
    return round(valor / EMU_POR_PT) if valor is not None else ""


# --------------------------------------------------------------------------
# Planilhas
# --------------------------------------------------------------------------

def criar_levantamento_vazio(caminho):
    """Planilha onde a estagiária cola as duas tabelas devolvidas pelo agente."""
    wb = Workbook()
    ws = wb.active
    ws.title = "moldes"
    ws.append(COLUNAS_MOLDES)
    ws2 = wb.create_sheet("caixas")
    ws2.append(COLUNAS_CAIXAS)
    for aba in (ws, ws2):
        for celula in aba[1]:
            celula.font = Font(bold=True)
        aba.freeze_panes = "A2"
    wb.save(caminho)


def main():
    parser = argparse.ArgumentParser(description="Extrai caixas e imagens dos slides.")
    parser.add_argument("arquivo", help="caminho do biblioteca.pptx")
    parser.add_argument("--saida", default="saida", help="pasta de saída (padrão: saida)")
    parser.add_argument("--sem-imagens", action="store_true",
                        help="não exporta os PNGs (útil para testar rápido)")
    args = parser.parse_args()

    arquivo = os.path.abspath(args.arquivo)
    if not os.path.exists(arquivo):
        sys.exit(f"Arquivo não encontrado: {arquivo}")

    pasta_saida = os.path.abspath(args.saida)
    pasta_img = os.path.join(pasta_saida, "miniaturas")
    pasta_txt = os.path.join(pasta_saida, "para_o_agente")
    os.makedirs(pasta_img, exist_ok=True)
    os.makedirs(pasta_txt, exist_ok=True)

    if not args.sem_imagens:
        print("Exportando imagens dos slides...")
        exportar_pngs(arquivo, pasta_img)

    prs = Presentation(arquivo)
    wb = Workbook()
    ws = wb.active
    ws.title = "caixas_extraidas"
    ws.append(["slide", "nome_atual", "grupo", "tipo", "texto_atual", "fonte_pt",
               "max_chars_estimado", "linhas_estimadas",
               "esquerda_pt", "topo_pt", "largura_pt", "altura_pt"])
    for celula in ws[1]:
        celula.font = Font(bold=True)
    ws.freeze_panes = "A2"

    for n, slide in enumerate(prs.slides, start=1):
        linhas_txt = [f"SLIDE {n:02d} (imagem: {n:02d}.png)", "Caixas:"]
        for shape, grupo in listar_caixas(slide.shapes):
            tipo = tipo_da_caixa(shape)
            texto = fonte = max_chars = linhas = ""
            if tipo == "texto":
                texto = shape.text_frame.text.replace("\n", " / ").strip()
                fonte = tamanho_fonte(shape)
                max_chars, linhas = estimar_limites(shape, fonte)

            ws.append([n, shape.name, grupo, tipo, texto, fonte, max_chars, linhas,
                       pt(shape.left), pt(shape.top), pt(shape.width), pt(shape.height)])

            detalhe = (f"- {shape.name} | tipo: {tipo} | posição: "
                       f"{pt(shape.left)},{pt(shape.top)} pt | tamanho: "
                       f"{pt(shape.width)}x{pt(shape.height)} pt")
            if grupo:
                detalhe += f" | dentro do grupo: {grupo}"
            if tipo == "texto":
                detalhe += (f" | máx {max_chars} caracteres, {linhas} linha(s)"
                            f" | texto: \"{texto[:80]}\"")
            linhas_txt.append(detalhe)

        with open(os.path.join(pasta_txt, f"{n:02d}.txt"), "w", encoding="utf-8") as f:
            f.write("\n".join(linhas_txt) + "\n")

    wb.save(os.path.join(pasta_saida, "caixas_extraidas.xlsx"))

    levantamento = os.path.join(pasta_saida, "levantamento.xlsx")
    if not os.path.exists(levantamento):  # nunca sobrescreve trabalho já feito
        criar_levantamento_vazio(levantamento)

    print(f"Pronto: {len(prs.slides)} slides processados em {pasta_saida}")


if __name__ == "__main__":
    main()
