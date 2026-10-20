"""
Agente 2 - Monta a apresentação a partir do JSON gerado pelo agente montador.

Para cada slide do pedido, copia o molde da biblioteca_final.pptx (mantendo o
design), preenche as caixas pelo nome e, no fim, remove os moldes que sobraram.

Uso:
  python montar_apresentacao.py pedido.json
  python montar_apresentacao.py pedido.json --biblioteca biblioteca_final.pptx \
      --levantamento levantamento.xlsx --saida minha_apresentacao.pptx

Formato dos campos aceitos no JSON:
  "titulo": "texto"                          -> caixa de texto (\\n quebra parágrafo)
  "texto_1": ["item 1", "item 2"]            -> um parágrafo por item
  "grafico_1": {"categorias": [...], "series": [{"nome": "...", "valores": [...]}]}
  "tabela_1": {"linhas": [["A", "B"], ["1", "2"]]}
  "imagem_1": "caminho/foto.png"             -> troca a imagem (se o arquivo existir)
  "imagem_1": "[sugestão: foto de equipe]"   -> mantém a imagem e anota a sugestão nas notas
"""

import argparse
import copy
import json
import os
import re
import sys
import unicodedata

from lxml import etree
from openpyxl import load_workbook
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.opc.packuri import PackURI
from pptx.oxml.ns import qn
from pptx.parts.chart import ChartPart
from pptx.parts.embeddedpackage import EmbeddedXlsxPart


# --------------------------------------------------------------------------
# Leitura do levantamento (moldes e limites das caixas)
# --------------------------------------------------------------------------

def _linhas_da_aba(wb, nome):
    if nome not in wb.sheetnames:
        return []
    linhas = list(wb[nome].iter_rows(values_only=True))
    if not linhas:
        return []
    cab = [str(c).strip().lower() if c else "" for c in linhas[0]]
    return [dict(zip(cab, l)) for l in linhas[1:] if l and any(v not in (None, "") for v in l)]


def ler_levantamento(caminho):
    """Devolve {molde: {"slide": n, "campos": {nome: (max_chars, linhas)}}}."""
    wb = load_workbook(caminho, read_only=True, data_only=True)
    moldes = {}
    por_slide = {}
    for l in _linhas_da_aba(wb, "moldes"):
        if not l.get("molde") or l.get("slide") in (None, ""):
            continue
        n = int(str(l["slide"]).strip())
        moldes[str(l["molde"]).strip()] = {"slide": n, "campos": {}}
        por_slide[n] = str(l["molde"]).strip()
    for l in _linhas_da_aba(wb, "caixas"):
        if l.get("slide") in (None, "") or not l.get("nome_novo"):
            continue
        molde = por_slide.get(int(str(l["slide"]).strip()))
        if molde:
            moldes[molde]["campos"][str(l["nome_novo"]).strip()] = (
                _int_ou_none(l.get("max_chars")), _int_ou_none(l.get("linhas")))
    return moldes


def _int_ou_none(valor):
    try:
        return int(valor)
    except (TypeError, ValueError):
        return None


# --------------------------------------------------------------------------
# Copiar um slide dentro da mesma apresentação
# --------------------------------------------------------------------------

def _novo_partname(package, modelo):
    """Gera um partname livre, ex.: /ppt/charts/chart%d.xml."""
    return package.next_partname(modelo)


def _clonar_grafico(slide_novo, chart_part_original):
    """Cria uma cópia independente do gráfico (e da planilha embutida dele)."""
    package = slide_novo.part.package
    novo = ChartPart.load(_novo_partname(package, "/ppt/charts/chart%d.xml"),
                          chart_part_original.content_type, package,
                          chart_part_original.blob)
    mapa = {}
    for rid, rel in chart_part_original.rels.items():
        if rel.is_external:
            mapa[rid] = novo.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.PACKAGE:  # planilha com os dados do gráfico
            mapa[rid] = novo.relate_to(EmbeddedXlsxPart.new(rel.target_part.blob, package),
                                       RT.PACKAGE)
        else:
            mapa[rid] = novo.relate_to(rel.target_part, rel.reltype)
    _trocar_rids(novo._element, mapa)
    return novo


R_NS = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}"


def _trocar_rids(elemento, mapa):
    for no in elemento.iter():
        for atributo, valor in list(no.attrib.items()):
            if atributo.startswith(R_NS) and valor in mapa:
                no.set(atributo, mapa[valor])


def duplicar_slide(prs, origem):
    """Adiciona ao fim da apresentação uma cópia fiel do slide de origem."""
    novo = prs.slides.add_slide(origem.slide_layout)
    arvore_nova = novo.shapes._spTree
    for el in list(arvore_nova):
        if el.tag.endswith("}sp") or el.tag.endswith("}pic") or el.tag.endswith("}graphicFrame") \
                or el.tag.endswith("}grpSp") or el.tag.endswith("}cxnSp"):
            arvore_nova.remove(el)

    # fundo próprio do slide, se houver
    bg = origem._element.cSld.find(qn("p:bg"))
    if bg is not None:
        novo._element.cSld.insert(0, copy.deepcopy(bg))

    # mapeia relações (imagens, gráficos, links) do slide de origem
    mapa_rid = {}
    for rid, rel in origem.part.rels.items():
        if rel.reltype in (RT.SLIDE_LAYOUT, RT.NOTES_SLIDE):
            continue
        if rel.is_external:
            novo_rid = novo.part.relate_to(rel.target_ref, rel.reltype, is_external=True)
        elif rel.reltype == RT.CHART:
            novo_rid = novo.part.relate_to(_clonar_grafico(novo, rel.target_part), RT.CHART)
        else:
            novo_rid = novo.part.relate_to(rel.target_part, rel.reltype)
        mapa_rid[rid] = novo_rid

    for el in origem.shapes._spTree:
        if el.tag.endswith("}nvGrpSpPr") or el.tag.endswith("}grpSpPr"):
            continue
        copia = copy.deepcopy(el)
        _trocar_rids(copia, mapa_rid)
        arvore_nova.append(copia)
    return novo


def remover_slide(prs, slide):
    lista = prs.slides._sldIdLst
    for sld_id in list(lista):
        if prs.part.related_part(sld_id.rId) is slide.part:
            prs.part.drop_rel(sld_id.rId)
            lista.remove(sld_id)
            return


# --------------------------------------------------------------------------
# Preencher as caixas
# --------------------------------------------------------------------------

def listar_caixas(shapes):
    for shape in shapes:
        yield shape
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            yield from listar_caixas(shape.shapes)


def escrever_texto(text_frame, valor):
    """Troca o texto mantendo a formatação do primeiro trecho de cada parágrafo."""
    paragrafos = valor if isinstance(valor, list) else str(valor).split("\n")
    paragrafos = [str(p) for p in paragrafos] or [""]
    txBody = text_frame._txBody
    existentes = txBody.findall(qn("a:p"))
    modelo = existentes[0]
    # usa o segundo parágrafo como modelo dos itens seguintes (ex.: tópicos), se houver
    modelo_resto = existentes[1] if len(existentes) > 1 else modelo
    for p in existentes:
        txBody.remove(p)

    for i, texto in enumerate(paragrafos):
        p = copy.deepcopy(modelo if i == 0 else modelo_resto)
        runs = p.findall(qn("a:r"))
        for extra in runs[1:]:
            p.remove(extra)
        for campo in p.findall(qn("a:fld")) + p.findall(qn("a:br")):
            p.remove(campo)
        if runs:
            runs[0].find(qn("a:t")).text = texto
        else:
            # caixa vazia no molde: cria o trecho com a formatação de fim de parágrafo
            r = etree.Element(qn("a:r"))
            end = p.find(qn("a:endParaRPr"))
            if end is not None:
                rpr = copy.deepcopy(end)
                rpr.tag = qn("a:rPr")
                r.append(rpr)
            etree.SubElement(r, qn("a:t")).text = texto
            if end is not None:
                end.addprevious(r)
            else:
                p.append(r)
        txBody.append(p)


def escrever_grafico(shape, valor):
    dados = CategoryChartData()
    dados.categories = valor.get("categorias", [])
    for serie in valor.get("series", []):
        dados.add_series(serie.get("nome", ""), serie.get("valores", []))
    shape.chart.replace_data(dados)


def escrever_tabela(shape, valor):
    linhas = valor.get("linhas", []) if isinstance(valor, dict) else valor
    tabela = shape.table
    for i, linha in enumerate(linhas[:len(tabela.rows)]):
        for j, texto in enumerate(linha[:len(tabela.columns)]):
            escrever_texto(tabela.cell(i, j).text_frame, texto)
    if len(linhas) > len(tabela.rows) or any(len(l) > len(tabela.columns) for l in linhas):
        return "a tabela do molde é menor que os dados; o excedente foi ignorado"
    return None


def trocar_imagem(slide, shape, caminho):
    _, rid = slide.part.get_or_add_image_part(caminho)
    blip = shape._element.find(".//" + qn("a:blip"))
    blip.set(qn("r:embed"), rid)


def contar_texto(valor):
    if isinstance(valor, list):
        return sum(len(str(v)) for v in valor), len(valor)
    texto = str(valor)
    return len(texto.replace("\n", "")), texto.count("\n") + 1


def preencher_slide(slide, campos, limites, pasta_base, avisos, rotulo):
    caixas = {c.name: c for c in listar_caixas(slide.shapes)}
    sugestoes = []
    for nome, valor in campos.items():
        if nome.startswith("fixo_"):
            avisos.append(f"{rotulo}: '{nome}' é decorativa e não foi alterada")
            continue
        shape = caixas.get(nome)
        if shape is None:
            avisos.append(f"{rotulo}: caixa '{nome}' não existe neste molde")
            continue

        if getattr(shape, "has_chart", False) and shape.has_chart:
            if isinstance(valor, dict):
                escrever_grafico(shape, valor)
            else:
                avisos.append(f"{rotulo}: '{nome}' é gráfico e precisa de categorias e series")
        elif getattr(shape, "has_table", False) and shape.has_table:
            erro = escrever_tabela(shape, valor)
            if erro:
                avisos.append(f"{rotulo}: '{nome}': {erro}")
        elif shape.shape_type == MSO_SHAPE_TYPE.PICTURE or nome.startswith("imagem_"):
            caminho = os.path.join(pasta_base, str(valor))
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and os.path.isfile(caminho):
                trocar_imagem(slide, shape, caminho)
            else:
                sugestoes.append(f"Trocar {nome}: {valor}")
        elif shape.has_text_frame:
            escrever_texto(shape.text_frame, valor)
            max_chars, max_linhas = limites.get(nome, (None, None))
            chars, linhas = contar_texto(valor)
            if max_chars and chars > max_chars:
                avisos.append(f"{rotulo}: '{nome}' tem {chars} caracteres (máx {max_chars})")
            if max_linhas and linhas > max_linhas:
                avisos.append(f"{rotulo}: '{nome}' tem {linhas} linhas (máx {max_linhas})")
        else:
            avisos.append(f"{rotulo}: '{nome}' não aceita conteúdo")

    vazias = [n for n in limites if n not in campos and not n.startswith("fixo_")]
    if vazias:
        avisos.append(f"{rotulo}: caixas sem conteúdo no pedido: {', '.join(vazias)}")
    return sugestoes


def escrever_notas(slide, notas, sugestoes):
    partes = [n for n in [notas] if n] + sugestoes
    if partes:
        slide.notes_slide.notes_text_frame.text = "\n".join(partes)


# --------------------------------------------------------------------------
# Programa principal
# --------------------------------------------------------------------------

def nome_de_arquivo(titulo):
    base = unicodedata.normalize("NFKD", titulo).encode("ascii", "ignore").decode()
    base = re.sub(r"[^A-Za-z0-9]+", "_", base).strip("_").lower()
    return (base or "apresentacao") + ".pptx"


def ler_pedido(caminho):
    with open(caminho, encoding="utf-8") as f:
        texto = f.read().strip()
    # aceita o JSON colado com a cerca ```json ... ``` do Copilot
    texto = re.sub(r"^```(?:json)?\s*|\s*```$", "", texto)
    try:
        return json.loads(texto)
    except json.JSONDecodeError as erro:
        sys.exit(f"O JSON do pedido está inválido (linha {erro.lineno}): {erro.msg}")


def main():
    parser = argparse.ArgumentParser(description="Monta uma apresentação a partir do pedido JSON.")
    parser.add_argument("pedido", help="arquivo JSON gerado pelo agente montador")
    parser.add_argument("--biblioteca", default="biblioteca_final.pptx")
    parser.add_argument("--levantamento", default="levantamento.xlsx")
    parser.add_argument("--saida", help="nome do arquivo final (padrão: baseado no título)")
    args = parser.parse_args()

    pedido = ler_pedido(args.pedido)
    moldes = ler_levantamento(args.levantamento) if os.path.exists(args.levantamento) else {}
    if not moldes:
        print("Aviso: levantamento não encontrado; usando só slide_biblioteca e sem checar limites.")

    prs = Presentation(args.biblioteca)
    originais = list(prs.slides)
    pasta_base = os.path.dirname(os.path.abspath(args.pedido))
    avisos = []

    for i, item in enumerate(pedido.get("slides", []), start=1):
        nome_molde = item.get("molde", "")
        info = moldes.get(nome_molde)
        numero = info["slide"] if info else item.get("slide_biblioteca")
        rotulo = f"slide {i} ({nome_molde or 'sem molde'})"
        if not numero or not 1 <= int(numero) <= len(originais):
            avisos.append(f"{rotulo}: molde não encontrado na biblioteca; slide pulado")
            continue
        if info and item.get("slide_biblioteca") and int(item["slide_biblioteca"]) != info["slide"]:
            avisos.append(f"{rotulo}: slide_biblioteca não bate com o levantamento; "
                          f"usado o slide {info['slide']}")

        novo = duplicar_slide(prs, originais[int(numero) - 1])
        limites = info["campos"] if info else {}
        sugestoes = preencher_slide(novo, item.get("campos", {}), limites,
                                    pasta_base, avisos, rotulo)
        escrever_notas(novo, item.get("notas", ""), sugestoes)

    for slide in originais:
        remover_slide(prs, slide)

    saida = args.saida or nome_de_arquivo(pedido.get("titulo_apresentacao", "apresentacao"))
    prs.save(saida)
    print(f"Apresentação salva em {saida} ({len(prs.slides)} slides)")
    if avisos:
        print("\nConfira:")
        for a in avisos:
            print(f"  - {a}")


if __name__ == "__main__":
    main()
