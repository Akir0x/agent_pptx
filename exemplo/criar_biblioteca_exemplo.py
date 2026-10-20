"""
Cria um biblioteca.pptx de exemplo (5 moldes) para testar o fluxo inteiro
sem precisar dos slides reais da empresa.

Uso:
  python criar_biblioteca_exemplo.py
"""

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

AZUL = RGBColor(0x1F, 0x3A, 0x5F)
AZUL_CLARO = RGBColor(0x3B, 0x82, 0xC4)
CINZA = RGBColor(0x60, 0x60, 0x60)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)


def caixa(slide, x, y, w, h, texto, tamanho, cor=AZUL, negrito=False):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = texto
    run.font.size = Pt(tamanho)
    run.font.bold = negrito
    run.font.color.rgb = cor
    return tb


def faixa(slide, x, y, w, h, cor):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = cor
    s.line.fill.background()
    return s


def main():
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    vazio = prs.slide_layouts[6]

    # 1. Capa
    s = prs.slides.add_slide(vazio)
    faixa(s, 0, 0, 0.4, 7.5, AZUL)
    caixa(s, 1, 2.4, 8, 1.4, "Título da apresentação", 40, negrito=True)
    caixa(s, 1, 3.9, 8, 0.6, "Subtítulo ou nome do cliente", 22, CINZA)
    caixa(s, 1, 5.8, 3, 0.5, "Mês 2026", 16, CINZA)

    # 2. Título + tópicos
    s = prs.slides.add_slide(vazio)
    caixa(s, 0.8, 0.5, 11, 0.9, "Título do slide", 32, negrito=True)
    tb = caixa(s, 0.8, 1.7, 11, 4.8, "Primeiro tópico", 20, CINZA)
    p = tb.text_frame.add_paragraph()
    r = p.add_run()
    r.text = "Segundo tópico"
    r.font.size, r.font.color.rgb = Pt(20), CINZA

    # 3. Gráfico de barras
    s = prs.slides.add_slide(vazio)
    caixa(s, 0.8, 0.5, 11, 0.9, "Título do gráfico", 32, negrito=True)
    dados = CategoryChartData()
    dados.categories = ["A", "B", "C"]
    dados.add_series("Série", (1, 2, 3))
    s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(1.6),
                       Inches(11.5), Inches(5.2), dados)

    # 4. Timeline com 3 etapas (cada etapa num grupo)
    s = prs.slides.add_slide(vazio)
    caixa(s, 0.8, 0.5, 11, 0.9, "Cronograma", 32, negrito=True)
    faixa(s, 0.8, 3.4, 11.7, 0.08, AZUL_CLARO)
    for i in range(3):
        x = 1.2 + i * 4
        grupo = s.shapes.add_group_shape()
        grupo.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(3.15), Inches(0.6), Inches(0.6))
        caixa(grupo, x - 0.2, 2.4, 3, 0.5, "Data", 14, AZUL_CLARO, True)
        caixa(grupo, x - 0.2, 4.0, 3.4, 0.5, "Nome da etapa", 18, AZUL, True)
        caixa(grupo, x - 0.2, 4.6, 3.4, 1.2, "Descrição curta da etapa", 14, CINZA)

    # 5. Destaque com número grande
    s = prs.slides.add_slide(vazio)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = AZUL
    caixa(s, 1, 2, 11.3, 2, "+00%", 96, BRANCO, True)
    caixa(s, 1, 4.3, 11.3, 1, "Frase explicando o número", 24, BRANCO)

    prs.save("biblioteca.pptx")
    print("biblioteca.pptx criado com 5 moldes")


if __name__ == "__main__":
    main()
